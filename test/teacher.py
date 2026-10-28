"""관제실에서 엑셀을 올려 문제 세트를 만드는 흐름을 실제 브라우저로 확인한다."""
import subprocess, time, os, signal, json, urllib.request
from playwright.sync_api import sync_playwright

DIR = '/home/claude/yut/english_yut_battle_v3'
PORT = 3160
PIN = 'test1234'
OUT = '/home/claude/yut/shots'
TMP = '/tmp/claude-0/-home-claude/26ac3424-1b12-527c-846d-d84ac0cdd7e9/scratchpad'
os.makedirs(OUT, exist_ok=True)

env = dict(os.environ, PORT=str(PORT), DATA_DIR='/tmp/yut-teacher-data', TEACHER_PIN=PIN)
srv = subprocess.Popen(['node', 'server.js'], cwd=DIR, env=env,
                       stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
time.sleep(1.6)
URL = f'http://localhost:{PORT}'
errors, results = [], []

def check(name, ok, detail=''):
    results.append((name, ok, detail))
    print(('  OK  ' if ok else '  실패 ') + name + (' — ' + detail if detail else ''))

try:
    # 1) 서버가 주는 양식을 그대로 내려받아 문제를 채운다 (선생님이 하는 그대로)
    tpl = f'{TMP}/양식.xlsx'
    urllib.request.urlretrieve(f'{URL}/api/template.xlsx', tpl)
    import openpyxl
    wb = openpyxl.load_workbook(tpl)
    ws = wb[wb.sheetnames[0]]
    hdr = [c.value for c in ws[1]]
    print('양식 첫 행:', hdr)
    # 기존 예시 줄을 지우고 새 문제 4개를 쓴다
    ws.delete_rows(2, ws.max_row)
    rows = [
        ['Lesson 10', 'What would you like to eat?', 'I would like a pizza.',
         'I am like a pizza.', 'I would liking pizza.', 'I like to pizza.', 1, 'easy', 'would like + 명사'],
        ['Lesson 10', 'How much is this cap?', "It's five dollars.", "It's five dollar.",
         "They're five dollars.", "It are five dollars.", 1, 'easy', '값을 말할 때 It is ~'],
        ['Lesson 10', 'Which is correct?', 'I will have bibimbap.', 'I will having bibimbap.',
         'I will has bibimbap.', 'I will be bibimbap.', 1, 'easy', ''],
        ['Lesson 10', 'What do you want for dessert?', 'Ice cream, please.', 'Ice cream, thank.',
         'Ice cream, sorry.', 'Ice cream, hello.', 1, 'easy', ''],
    ]
    for r in rows:
        ws.append(r)
    xl = f'{TMP}/10단원_문제.xlsx'
    wb.save(xl)
    check('엑셀 양식 내려받기 + 문제 채우기', True, f'{len(rows)}문제')

    with sync_playwright() as pw:
        br = pw.chromium.launch()
        pg = br.new_context(viewport={'width': 1280, 'height': 880}).new_page()
        pg.on('pageerror', lambda e: errors.append(str(e)))
        pg.on('console', lambda m: errors.append('console: ' + m.text) if m.type == 'error' else None)
        pg.goto(URL)
        pg.evaluate("() => document.querySelectorAll('details.moreBox').forEach(d=>d.open=true)")

        # 2) 관제실 열기
        pg.click('details.teacherBox summary')
        pg.fill('#teacherPin', PIN)
        pg.click('#btnTeacher')
        pg.wait_for_selector('#teacher.on', timeout=8000)
        check('관제실 열기', True)

        # 3) 새 문제 세트 + 엑셀 업로드
        pg.click('#btnNewSet')
        pg.fill('#className', '6-1 10단원 음식')
        pg.set_input_files('#classFile', xl)
        pg.click('#btnMakeClass')
        pg.wait_for_selector('#classOut .hint', timeout=12000)
        out = pg.inner_text('#classOut')
        toast = pg.inner_text('#toast')
        ok = '만들었습니다' in out
        check('엑셀 업로드로 문제 세트 만들기', ok, out.replace('\n', ' ')[:110] or toast)
        pg.screenshot(path=f'{OUT}/13_문제세트_업로드.png')

        if ok:
            # 4) 만들어진 주소로 문제가 실제로 들어갔는지
            import re
            m = re.search(r'\?class=([A-Z0-9]+)', out)
            code = m.group(1) if m else None
            if code:
                with urllib.request.urlopen(f'{URL}/api/class/{code}') as r:
                    d = json.load(r)
                check('학생 주소에서 문제 세트 확인', d.get('count') == len(rows),
                      f"{d.get('name')} · 문제 {d.get('count')}개")

            # 5) 지금 적용
            pg.click(f'[data-act="{code}"]')
            time.sleep(0.8)
            t = pg.inner_text('#toast')
            check('지금 적용', '적용했습니다' in t, t)

            # 6) 학생이 기본 주소로 방을 만들면 그 문제가 나오는지
            s = br.new_context(viewport={'width': 1100, 'height': 800}).new_page()
            s.goto(URL)
            s.evaluate("() => document.querySelectorAll('details.moreBox').forEach(d=>d.open=true)")
            s.fill('#myName', '학생1')
            s.click('#btnCreate')
            s.wait_for_selector('#waiting.on', timeout=8000)
            code2 = s.inner_text('#waitCode').strip()
            s2 = br.new_context(viewport={'width': 1100, 'height': 800}).new_page()
            s2.goto(URL)
            s2.evaluate("() => document.querySelectorAll('details.moreBox').forEach(d=>d.open=true)")
            s2.fill('#myName', '학생2')
            s2.fill('#roomCode', code2)
            s2.click('#btnJoin')
            s2.wait_for_selector('#waiting.on', timeout=8000)
            s.click('#btnStart')
            s.wait_for_selector('#ovBattle:not([hidden])', timeout=10000)
            q = s.inner_text('#qText')
            unit = s.inner_text('#qUnit')
            check('학생 방에 올린 문제가 나오는지', 'Lesson 10' in unit, f'{unit} / {q}')
            s.screenshot(path=f'{OUT}/14_올린문제_출제.png')

            # 7) 세트 삭제
            pg.once('dialog', lambda d: d.accept())
            pg.click(f'[data-del="{code}"]')
            time.sleep(0.8)
            check('문제 세트 삭제', '지웠습니다' in pg.inner_text('#toast'), pg.inner_text('#toast'))

        # 8) 틀린 PIN은 거부되는지
        pg.evaluate("() => { teacherPin = '0000'; }")
        pg.click('#btnNewSet')
        pg.fill('#className', '잘못된 핀 시험')
        pg.click('#btnMakeClass')
        time.sleep(1.0)
        t = pg.inner_text('#toast')
        check('틀린 PIN은 막히는지', 'PIN' in t, t)

        br.close()
finally:
    srv.send_signal(signal.SIGTERM); time.sleep(0.4); srv.kill()
    err = srv.stderr.read().decode()[-800:] if srv.stderr else ''

print('\n' + '─' * 50)
bad = [r for r in results if not r[1]]
print(f'{len(results) - len(bad)} / {len(results)} 통과')
real = [e for e in errors if 'favicon' not in e and 'TUNNEL' not in e]
print('브라우저 오류:', sorted(set(real))[:5] if real else '없음')
if err.strip():
    print('서버 stderr:', err[-400:])
